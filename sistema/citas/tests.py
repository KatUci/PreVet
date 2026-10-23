"""Pruebas de las citas y del modelo de puntuación del riesgo.

Django levanta una base de datos aparte para correrlas, así que no tocan
los datos reales del sistema.
"""

from datetime import datetime, timedelta

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone

from mascotas.models import Cliente, Especie, Mascota, Tenencia
from .models import Cita, TipoAtencion
from . import riesgo


class BaseCitas(TestCase):
    """Arma una clínica mínima que sirve de punto de partida a las pruebas."""

    def setUp(self):
        self.vet = User.objects.create_user('vet', first_name='Paz',
                                            last_name='Herrera')
        self.tutor = User.objects.create_user('tutor', first_name='Ana',
                                              last_name='Soto')
        self.especie = Especie.objects.create(nombre='Perro')
        self.mascota = Mascota.objects.create(
            nombre='Bruno', especie=self.especie, sexo='macho')
        self.cliente = Cliente.objects.create(
            rut='15234871-1', nombres='Ana', apellidos='Soto')
        Tenencia.objects.create(cliente=self.cliente, mascota=self.mascota,
                                es_principal=True)
        self.control = TipoAtencion.objects.create(
            nombre='Control', duracion_minutos=30, precio=20000)

        # El modelo no permite dos citas del mismo veterinario a la misma
        # fecha y hora, así que cada una recibe un bloque distinto.
        self._bloque = 0

    def crear_cita(self, estado='pendiente', dias=7, tipo=None, mascota=None,
                   anticipacion=None):
        """Crea una cita y, si se pide, le fija los días de anticipación.

        fecha_creacion es auto_now_add, de modo que al crear queda siempre en
        hoy. Para probar la anticipación hay que sobrescribirla después.
        """
        self._bloque += 1
        hora = (datetime.min + timedelta(hours=8, minutes=5 * self._bloque)).time()

        mascota = mascota or self.mascota
        cita = Cita.objects.create(
            veterinario=self.vet,
            mascota=mascota,
            cliente=mascota.responsable,
            tipo_atencion=tipo or self.control,
            fecha=timezone.localdate() + timedelta(days=dias),
            hora=hora,
            estado=estado,
        )
        if anticipacion is not None:
            creada = cita.fecha - timedelta(days=anticipacion)
            Cita.objects.filter(pk=cita.pk).update(
                fecha_creacion=timezone.make_aware(
                    datetime.combine(creada, datetime.min.time())))
            cita.refresh_from_db()
        return cita


class CitaTest(BaseCitas):
    """El modelo Cita y lo que aporta este proyecto sobre el sistema base."""

    def test_el_estado_no_asistio_existe(self):
        """Es el aporte central: distinguir una falta de una cancelación."""
        estados = dict(Cita.ESTADOS)
        self.assertIn('no_asistio', estados)
        self.assertIn('cancelada', estados)
        self.assertNotEqual(estados['no_asistio'], estados['cancelada'])

    def test_valor_perdido_solo_cuando_el_tutor_no_se_presenta(self):
        self.assertEqual(self.crear_cita(estado='no_asistio').valor_perdido,
                         self.control.precio)
        self.assertEqual(self.crear_cita(estado='completada').valor_perdido, 0)
        self.assertEqual(self.crear_cita(estado='cancelada').valor_perdido, 0)

    def test_dias_anticipacion(self):
        self.assertEqual(self.crear_cita(anticipacion=10).dias_anticipacion, 10)
        self.assertEqual(self.crear_cita(anticipacion=0).dias_anticipacion, 0)


class EstadisticasTest(BaseCitas):
    """Las tasas históricas que alimentan el modelo."""

    def test_sin_historial_la_tasa_global_es_cero(self):
        self.assertEqual(riesgo.estadisticas()['global'], 0)

    def test_la_tasa_global_considera_solo_citas_resueltas(self):
        for _ in range(3):
            self.crear_cita(estado='completada')
        self.crear_cita(estado='no_asistio')
        self.crear_cita(estado='pendiente')   # no debería contar
        self.crear_cita(estado='cancelada')   # tampoco

        self.assertAlmostEqual(riesgo.estadisticas()['global'], 0.25)

    def test_registra_el_historial_por_tutor(self):
        self.crear_cita(estado='no_asistio')
        self.crear_cita(estado='completada')

        faltas, total = riesgo.estadisticas()['tutor'][self.cliente.pk]
        self.assertEqual((faltas, total), (1, 2))


class PuntuarTest(BaseCitas):
    """El corazón del proyecto: cómo se arma el puntaje de riesgo."""

    def test_los_pesos_suman_cien(self):
        self.assertEqual(
            riesgo.PESO_TUTOR + riesgo.PESO_ANTICIPACION
            + riesgo.PESO_DIA + riesgo.PESO_TIPO, 100)

    def test_el_puntaje_nunca_supera_cien(self):
        for _ in range(6):
            self.crear_cita(estado='no_asistio', dias=-3)
        cita = self.crear_cita(anticipacion=60)
        puntaje, _, _, _ = riesgo.puntuar(cita, riesgo.estadisticas())
        self.assertLessEqual(puntaje, 100)

    def test_quien_falta_siempre_puntua_mas_que_quien_nunca_falta(self):
        """Es el supuesto principal del modelo y debe quedar demostrado."""
        cumplidor = Cliente.objects.create(rut='16890342-1', nombres='Matías',
                                           apellidos='Silva')
        otra = Mascota.objects.create(nombre='Kira', especie=self.especie,
                                      sexo='hembra')
        Tenencia.objects.create(cliente=cumplidor, mascota=otra, es_principal=True)

        for _ in range(4):
            self.crear_cita(estado='no_asistio', dias=-5)
            self.crear_cita(estado='completada', dias=-5, mascota=otra)

        stats = riesgo.estadisticas()
        riesgoso = riesgo.puntuar(self.crear_cita(anticipacion=7), stats)[0]
        confiable = riesgo.puntuar(
            self.crear_cita(anticipacion=7, mascota=otra), stats)[0]

        self.assertGreater(riesgoso, confiable)

    def test_el_tutor_nuevo_se_trata_como_incertidumbre(self):
        """Sin historial no se asume inocencia: se asume que no se sabe."""
        cita = self.crear_cita(anticipacion=0)
        _, _, _, detalle = riesgo.puntuar(cita, riesgo.estadisticas())
        self.assertIn('Primera visita, sin historial', detalle)

    def test_el_detalle_explica_los_cuatro_factores(self):
        """Un puntaje sin explicación no le sirve a la recepcionista."""
        cita = self.crear_cita(anticipacion=5)
        _, _, _, detalle = riesgo.puntuar(cita, riesgo.estadisticas())
        self.assertEqual(len(detalle), 4)

    def test_cada_nivel_trae_su_accion_recomendada(self):
        cita = self.crear_cita(anticipacion=5)
        _, nivel, accion, _ = riesgo.puntuar(cita, riesgo.estadisticas())
        self.assertIn(nivel, ('alto', 'medio', 'bajo'))
        self.assertEqual(accion, riesgo.ACCIONES[nivel])

    def test_los_umbrales_estan_ordenados(self):
        self.assertGreater(riesgo.UMBRAL_ALTO, riesgo.UMBRAL_MEDIO)


class VistasCitasTest(BaseCitas):
    """La pantalla de citas y la acción de registrar una inasistencia."""

    def setUp(self):
        super().setUp()
        # Las pantallas exigen permisos: el usuario de prueba entra como Administración.
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Administración'))
        self.client.force_login(self.tutor)

    def test_el_listado_exige_haber_iniciado_sesion(self):
        self.client.logout()
        respuesta = self.client.get('/citas/')
        self.assertEqual(respuesta.status_code, 302)
        self.assertIn('/login/', respuesta.url)

    def test_el_listado_muestra_las_citas(self):
        self.crear_cita()
        respuesta = self.client.get('/citas/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, 'Bruno')

    def test_el_filtro_deja_solo_las_citas_de_ese_estado(self):
        self.crear_cita(estado='no_asistio')
        self.crear_cita(estado='completada')
        self.crear_cita(estado='completada')

        respuesta = self.client.get('/citas/?estado=no_asistio')
        self.assertEqual(respuesta.context['total'], 1)

    def test_cada_filtro_trae_su_propio_conteo(self):
        self.crear_cita(estado='no_asistio')
        self.crear_cita(estado='completada')

        respuesta = self.client.get('/citas/')
        conteos = {e['valor']: e['conteo'] for e in respuesta.context['estados']}
        self.assertEqual(conteos['no_asistio'], 1)
        self.assertEqual(conteos['completada'], 1)
        self.assertEqual(conteos['cancelada'], 0)

    def test_marcar_no_asistio_cambia_el_estado(self):
        # Solo se puede faltar a una cita que ya ocurrió: se usa una de ayer.
        cita = self.crear_cita(estado='pendiente', dias=-1)
        self.client.post(f'/citas/{cita.pk}/no-asistio/')
        cita.refresh_from_db()

        self.assertEqual(cita.estado, 'no_asistio')
        self.assertIsNotNone(cita.fecha_marcada_no_asistio)

    def test_marcar_no_asistio_no_acepta_una_visita_normal(self):
        """Una acción que cambia datos no puede ocurrir al abrir un enlace."""
        cita = self.crear_cita(estado='pendiente')
        self.client.get(f'/citas/{cita.pk}/no-asistio/')
        cita.refresh_from_db()

        self.assertEqual(cita.estado, 'pendiente')


class ClienteDeLaCitaTest(BaseCitas):
    """La cita solo puede quedar a nombre de un cliente vinculado a la mascota."""

    def test_cliente_ajeno_a_la_mascota_no_valida(self):
        from django.core.exceptions import ValidationError
        ajeno = Cliente.objects.create(rut='16890342-1', nombres='Matías', apellidos='Silva')
        cita = Cita(veterinario=self.vet, mascota=self.mascota, cliente=ajeno,
                    tipo_atencion=self.control, fecha=timezone.localdate(),
                    hora=datetime.min.time())
        with self.assertRaises(ValidationError):
            cita.full_clean()


class AgendarTest(TestCase):
    """El flujo completo: RUT del cliente, una de sus mascotas y la hora."""

    CLAVE = 'Huellitas.2026!'

    def setUp(self):
        from django.contrib.auth.models import Group
        from usuarios.models import Perfil
        from .models import HorarioAtencion

        self.vet = User.objects.create_user('vet', first_name='Paz', last_name='Herrera')
        Perfil.objects.create(usuario=self.vet, rol='veterinario')
        # Atiende de lunes a viernes, de 09:00 a 13:00.
        for dia in range(5):
            HorarioAtencion.objects.create(veterinario=self.vet, dia_semana=dia,
                                           hora_inicio='09:00', hora_fin='13:00')

        self.recepcion = User.objects.create_user('recepcion', password=self.CLAVE)
        self.recepcion.groups.add(Group.objects.get(name='Recepción'))

        canino = Especie.objects.create(nombre='Perro')
        self.ana = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        self.bruno = Mascota.objects.create(nombre='Bruno', especie=canino, sexo='macho')
        Tenencia.objects.create(cliente=self.ana, mascota=self.bruno, es_principal=True)
        self.otro = Cliente.objects.create(rut='16890342-1', nombres='Matías', apellidos='Silva')
        self.ajena = Mascota.objects.create(nombre='Kira', especie=canino, sexo='hembra')
        Tenencia.objects.create(cliente=self.otro, mascota=self.ajena, es_principal=True)

        self.control = TipoAtencion.objects.create(nombre='Control', duracion_minutos=30, precio=20000)
        self.client.login(username='recepcion', password=self.CLAVE)

        # El próximo lunes, para que siempre sea un día con horario.
        hoy = timezone.localdate()
        self.lunes = hoy + timedelta(days=7 - hoy.weekday())

    def url(self, rut='15.234.871-1'):
        return f'/citas/nueva/?rut={rut}'

    def datos(self, **cambios):
        datos = {'mascota': self.bruno.pk, 'tipo_atencion': self.control.pk,
                 'veterinario': self.vet.pk, 'fecha': self.lunes.isoformat(),
                 'hora': '10:00', 'notas': ''}
        datos.update(cambios)
        return datos

    def test_sin_permiso_de_agendar_recibe_403(self):
        User.objects.create_user('sinpermiso', password=self.CLAVE)
        self.client.login(username='sinpermiso', password=self.CLAVE)
        self.assertEqual(self.client.get('/citas/nueva/').status_code, 403)

    def test_rut_invalido(self):
        respuesta = self.client.get(self.url('12.345.678-9'))
        self.assertContains(respuesta, 'dígito verificador no corresponde')

    def test_rut_sin_cliente_ofrece_registrarlo(self):
        respuesta = self.client.get(self.url('12.345.678-5'))
        self.assertContains(respuesta, 'Registrar cliente')

    def test_solo_aparecen_las_mascotas_del_cliente(self):
        respuesta = self.client.get(self.url())
        self.assertContains(respuesta, 'Bruno')
        self.assertNotContains(respuesta, 'Kira')

    def test_agenda_la_hora_a_nombre_del_cliente(self):
        respuesta = self.client.post(self.url(), self.datos())
        self.assertRedirects(respuesta, '/citas/')
        cita = Cita.objects.get()
        self.assertEqual(cita.cliente, self.ana)
        self.assertEqual(cita.mascota, self.bruno)
        self.assertEqual(cita.estado, 'pendiente')

    def test_no_acepta_una_mascota_de_otro_cliente(self):
        self.client.post(self.url(), self.datos(mascota=self.ajena.pk))
        self.assertFalse(Cita.objects.exists())

    def test_no_acepta_fechas_pasadas(self):
        ayer = timezone.localdate() - timedelta(days=1)
        respuesta = self.client.post(self.url(), self.datos(fecha=ayer.isoformat()))
        self.assertContains(respuesta, 'ya pasó')
        self.assertFalse(Cita.objects.exists())

    def test_fuera_del_horario_del_profesional(self):
        respuesta = self.client.post(self.url(), self.datos(hora='15:00'))
        self.assertContains(respuesta, 'atiende ese día de 09:00 a 13:00')
        self.assertFalse(Cita.objects.exists())

    def test_no_permite_tomar_una_hora_ocupada(self):
        self.client.post(self.url(), self.datos())
        respuesta = self.client.post(self.url(), self.datos())
        self.assertContains(respuesta, 'ya está tomada')
        self.assertEqual(Cita.objects.count(), 1)


class RegistrarClienteYMascotaTest(TestCase):

    CLAVE = 'Huellitas.2026!'

    def setUp(self):
        from django.contrib.auth.models import Group
        recepcion = User.objects.create_user('recepcion', password=self.CLAVE)
        recepcion.groups.add(Group.objects.get(name='Recepción'))
        self.client.login(username='recepcion', password=self.CLAVE)

    def test_registra_cliente_con_rut_limpio_y_vuelve_a_agendar(self):
        respuesta = self.client.post('/clientes/nuevo/', {
            'rut': '12.345.678-5', 'nombres': 'Ana', 'apellidos': 'Soto'})
        self.assertRedirects(respuesta, '/citas/nueva/?rut=12345678-5', fetch_redirect_response=False)
        self.assertTrue(Cliente.objects.filter(rut='12345678-5').exists())

    def test_no_permite_rut_repetido(self):
        Cliente.objects.create(rut='12345678-5', nombres='Ana', apellidos='Soto')
        respuesta = self.client.post('/clientes/nuevo/', {
            'rut': '12345678-5', 'nombres': 'Otra', 'apellidos': 'Persona'})
        self.assertContains(respuesta, 'Ya hay un cliente registrado con este RUT.')

    def test_no_redirige_a_otros_sitios(self):
        respuesta = self.client.post('/clientes/nuevo/?next=https://sitio-malo.com/', {
            'rut': '12.345.678-5', 'nombres': 'Ana', 'apellidos': 'Soto'})
        self.assertNotIn('sitio-malo', respuesta['Location'])

    def test_registra_mascota_con_tipo_y_raza_y_la_vincula(self):
        from mascotas.models import Raza
        ana = Cliente.objects.create(rut='12345678-5', nombres='Ana', apellidos='Soto')
        ave = Especie.objects.get(nombre='Ave')
        loro, _ = Raza.objects.get_or_create(especie=ave, nombre='Loro')
        self.client.post(f'/clientes/{ana.pk}/mascotas/nueva/', {
            'especie': ave.pk, 'raza': loro.pk, 'nombre': 'Pepe', 'sexo': 'macho',
            'fecha_nacimiento': '', 'relacion': 'dueno'})
        pepe = Mascota.objects.get(nombre='Pepe')
        self.assertEqual(pepe.especie, ave)
        self.assertEqual(pepe.responsable, ana)

    def test_raza_de_otro_tipo_no_valida(self):
        from mascotas.models import Raza
        ana = Cliente.objects.create(rut='12345678-5', nombres='Ana', apellidos='Soto')
        felino = Especie.objects.get(nombre='Felino')
        loro, _ = Raza.objects.get_or_create(especie=Especie.objects.get(nombre='Ave'), nombre='Loro')
        respuesta = self.client.post(f'/clientes/{ana.pk}/mascotas/nueva/', {
            'especie': felino.pk, 'raza': loro.pk, 'nombre': 'Michi', 'sexo': 'macho',
            'fecha_nacimiento': '', 'relacion': 'dueno'})
        self.assertContains(respuesta, 'no es una raza de felino')
        self.assertFalse(Mascota.objects.filter(nombre='Michi').exists())


class CalendarioTest(TestCase):
    """El calendario y lo que se puede hacer con cada cita desde él."""

    CLAVE = 'Huellitas.2026!'

    def setUp(self):
        from django.contrib.auth.models import Group
        from usuarios.models import Perfil
        from .models import HorarioAtencion

        self.vet = User.objects.create_user('vet', first_name='Paz', last_name='Herrera')
        Perfil.objects.create(usuario=self.vet, rol='veterinario')
        self.otro_vet = User.objects.create_user('vet2', first_name='Luis', last_name='Contreras')
        Perfil.objects.create(usuario=self.otro_vet, rol='veterinario')
        for v in (self.vet, self.otro_vet):
            for dia in range(7):
                HorarioAtencion.objects.create(veterinario=v, dia_semana=dia,
                                               hora_inicio='09:00', hora_fin='13:00')

        self.recepcion = User.objects.create_user('recepcion', password=self.CLAVE)
        self.recepcion.groups.add(Group.objects.get(name='Recepción'))

        canino = Especie.objects.create(nombre='Perro')
        self.ana = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        self.bruno = Mascota.objects.create(nombre='Bruno', especie=canino, sexo='macho')
        Tenencia.objects.create(cliente=self.ana, mascota=self.bruno, es_principal=True)
        self.control = TipoAtencion.objects.create(nombre='Control', duracion_minutos=30, precio=20000)

        self.hoy = timezone.localdate()
        self.manana = self.hoy + timedelta(days=1)
        self.client.login(username='recepcion', password=self.CLAVE)

    def cita(self, dias=1, hora='10:00', vet=None, estado='pendiente'):
        return Cita.objects.create(
            veterinario=vet or self.vet, mascota=self.bruno, cliente=self.ana,
            tipo_atencion=self.control, fecha=self.hoy + timedelta(days=dias),
            hora=hora, estado=estado)

    def eventos(self, **params):
        desde = self.hoy - timedelta(days=3)
        hasta = self.hoy + timedelta(days=7)
        query = {'start': desde.isoformat(), 'end': hasta.isoformat(), **params}
        return self.client.get('/citas/calendario/eventos/', query).json()

    def test_la_pagina_carga(self):
        self.assertContains(self.client.get('/citas/calendario/'), 'id="calendario"')

    def test_entrega_las_citas_con_inicio_y_termino(self):
        self.cita()
        evento = self.eventos()[0]
        self.assertTrue(evento['start'].endswith('10:00:00'))
        self.assertTrue(evento['end'].endswith('10:30:00'))   # dura 30 minutos
        self.assertIn('Bruno', evento['title'])

    def test_filtra_por_profesional(self):
        self.cita(vet=self.vet)
        self.cita(vet=self.otro_vet, hora='11:00')
        self.assertEqual(len(self.eventos(veterinario=self.otro_vet.pk)), 1)

    def test_rango_invalido(self):
        respuesta = self.client.get('/citas/calendario/eventos/', {'start': 'x', 'end': 'y'})
        self.assertEqual(respuesta.status_code, 400)

    def test_el_detalle_muestra_al_cliente_y_el_riesgo(self):
        c = self.cita()
        respuesta = self.client.get(f'/citas/{c.pk}/detalle/')
        self.assertContains(respuesta, 'Ana Soto')
        self.assertContains(respuesta, 'Riesgo')

    def test_cancelar_guarda_motivo_y_libera_la_hora(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/cancelar/', {'motivo': 'Otro', 'detalle': 'Viaje'})
        c.refresh_from_db()
        self.assertEqual(c.estado, 'cancelada')
        self.assertEqual(c.motivo_cancelacion, 'Otro: Viaje')
        self.assertEqual(c.cancelada_por, self.recepcion)
        # El mismo bloque ya se puede volver a ocupar.
        self.cita()
        self.assertEqual(Cita.objects.filter(fecha=self.manana, hora='10:00').count(), 2)

    def test_cancelar_sin_motivo_no_hace_nada(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/cancelar/', {'motivo': ''})
        c.refresh_from_db()
        self.assertEqual(c.estado, 'pendiente')

    def test_reagendar_respeta_el_horario_y_las_horas_tomadas(self):
        c = self.cita()
        self.cita(hora='11:00')
        fuera = self.client.post(f'/citas/{c.pk}/reagendar/', {
            'veterinario': self.vet.pk, 'fecha': self.manana.isoformat(), 'hora': '15:00'})
        self.assertContains(fuera, 'atiende ese día')
        tomada = self.client.post(f'/citas/{c.pk}/reagendar/', {
            'veterinario': self.vet.pk, 'fecha': self.manana.isoformat(), 'hora': '11:00'})
        self.assertContains(tomada, 'ya está tomada')
        self.client.post(f'/citas/{c.pk}/reagendar/', {
            'veterinario': self.otro_vet.pk, 'fecha': self.manana.isoformat(), 'hora': '12:00'})
        c.refresh_from_db()
        self.assertEqual((c.veterinario, c.hora.strftime('%H:%M')), (self.otro_vet, '12:00'))

    def test_llegada(self):
        c = self.cita(dias=0)
        self.client.post(f'/citas/{c.pk}/llego/')
        c.refresh_from_db()
        self.assertIsNotNone(c.hora_llegada)
        self.assertEqual(c.estado, 'pendiente')

    def test_recepcion_no_cierra_la_atencion(self):
        """Cerrar la atención pide datos clínicos: solo lo hacen los veterinarios."""
        c = self.cita(dias=0)
        respuesta = self.client.post(f'/citas/{c.pk}/atender/', {'motivo': 'Control', 'diagnostico': 'Sano'})
        self.assertEqual(respuesta.status_code, 403)
        c.refresh_from_db()
        self.assertEqual(c.estado, 'pendiente')

    def test_una_cita_cancelada_no_se_marca_como_inasistencia(self):
        c = self.cita(dias=0, estado='cancelada')
        self.client.post(f'/citas/{c.pk}/no-asistio/')
        c.refresh_from_db()
        self.assertEqual(c.estado, 'cancelada')

    def test_las_acciones_exigen_poder_modificar_citas(self):
        c = self.cita()
        User.objects.create_user('mirón', password=self.CLAVE)
        self.client.login(username='mirón', password=self.CLAVE)
        self.assertEqual(self.client.post(f'/citas/{c.pk}/cancelar/', {'motivo': 'Otro'}).status_code, 403)


class WhatsAppTest(CalendarioTest):
    """El panel de la cita trae un enlace de WhatsApp con el recordatorio."""

    def test_enlace_con_codigo_de_chile_y_mensaje(self):
        self.ana.telefono = '9 1234 5678'
        self.ana.save()
        respuesta = self.client.get(f'/citas/{self.cita().pk}/detalle/')
        self.assertContains(respuesta, 'https://wa.me/56912345678?text=Hola%20Ana')

    def test_sin_telefono_no_hay_enlace(self):
        respuesta = self.client.get(f'/citas/{self.cita().pk}/detalle/')
        self.assertNotContains(respuesta, 'wa.me')
        self.assertContains(respuesta, 'Sin teléfono registrado')


class ListadoProfesionalTest(BaseCitas):
    """Buscador, filtros, paginador y exportación del listado de citas."""

    def setUp(self):
        super().setUp()
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Administración'))
        self.client.force_login(self.tutor)

    def test_busca_por_nombre_de_la_mascota(self):
        self.crear_cita()
        otra = Mascota.objects.create(nombre='Luna', especie=self.especie, sexo='hembra')
        Tenencia.objects.create(cliente=self.cliente, mascota=otra, es_principal=True)
        self.crear_cita(mascota=otra)
        respuesta = self.client.get('/citas/?q=luna')
        self.assertEqual(respuesta.context['total'], 1)

    def test_busca_por_rut_con_o_sin_puntos(self):
        self.crear_cita()
        for texto in ('15.234.871-1', '15234871', '152348711'):
            with self.subTest(texto=texto):
                self.assertEqual(self.client.get(f'/citas/?q={texto}').context['total'], 1)
        self.assertEqual(self.client.get('/citas/?q=99999999').context['total'], 0)

    def test_filtra_por_rango_de_fechas(self):
        self.crear_cita(dias=-3)
        self.crear_cita(dias=0)
        self.crear_cita(dias=5)
        hoy = timezone.localdate()
        respuesta = self.client.get(f'/citas/?desde={hoy:%Y-%m-%d}&hasta={hoy:%Y-%m-%d}')
        self.assertEqual(respuesta.context['total'], 1)

    def test_una_fecha_mal_escrita_no_rompe_la_pantalla(self):
        self.crear_cita()
        respuesta = self.client.get('/citas/?desde=ayer')
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['filtros']['desde'], '')

    def test_pagina_de_a_veinte_y_conserva_los_filtros(self):
        for _ in range(25):
            self.crear_cita(estado='completada')
        respuesta = self.client.get('/citas/?estado=completada')
        self.assertEqual(len(respuesta.context['pagina'].object_list), 20)
        self.assertContains(respuesta, 'estado=completada&pagina=2')
        segunda = self.client.get('/citas/?estado=completada&pagina=2')
        self.assertEqual(len(segunda.context['pagina'].object_list), 5)

    def test_exporta_a_excel_las_citas_filtradas(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.crear_cita(estado='no_asistio')
        self.crear_cita(estado='completada')
        respuesta = self.client.get('/citas/exportar/?estado=no_asistio')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        hoja = load_workbook(BytesIO(respuesta.content)).active
        self.assertEqual(hoja.max_row, 2)          # encabezado + 1 cita
        self.assertEqual(hoja['C2'].value, 'Bruno')
        self.assertEqual(hoja['J2'].value, 'No asistió')


class PanelDelDiaTest(BaseCitas):
    """El inicio de recepción muestra el flujo de pacientes del día."""

    def setUp(self):
        super().setUp()
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Recepción'))
        self.client.force_login(self.tutor)

    def test_recepcion_ve_el_panel_en_el_inicio(self):
        respuesta = self.client.get('/')
        self.assertTemplateUsed(respuesta, 'citas/panel_dia.html')

    def test_administracion_sigue_viendo_los_indicadores(self):
        from django.contrib.auth.models import Group
        self.tutor.groups.set([Group.objects.get(name='Administración')])
        respuesta = self.client.get('/')
        self.assertTemplateUsed(respuesta, 'core/inicio.html')
        self.assertEqual(self.client.get('/citas/hoy/').status_code, 200)

    def test_reparte_las_citas_de_hoy_por_etapa(self):
        por_llegar = self.crear_cita(dias=0)
        en_sala = self.crear_cita(dias=0)
        en_sala.hora_llegada = timezone.now()
        en_sala.save()
        self.crear_cita(dias=0, estado='completada')
        self.crear_cita(dias=0, estado='cancelada')

        contexto = self.client.get('/').context
        self.assertEqual([t['cita'] for t in contexto['por_llegar']], [por_llegar])
        self.assertEqual([t['cita'] for t in contexto['en_sala']], [en_sala])
        self.assertEqual(len(contexto['atendidas']), 1)
        self.assertEqual(contexto['total_hoy'], 3)    # la cancelada no cuenta

    def test_marca_al_tutor_que_ya_falto_varias_veces(self):
        self.crear_cita(dias=-10, estado='no_asistio')
        self.crear_cita(dias=-5, estado='no_asistio')
        self.crear_cita(dias=0)
        tarjeta = self.client.get('/').context['por_llegar'][0]
        self.assertEqual(tarjeta['faltas'], 2)

    def test_lista_las_citas_de_manana_para_confirmar(self):
        self.crear_cita(dias=1)
        self.crear_cita(dias=2)
        contexto = self.client.get('/').context
        self.assertEqual(contexto['confirmar_total'], 1)


class RiesgoPantallaTest(BaseCitas):
    """Filtros, resumen y exportación de Horas en riesgo."""

    def setUp(self):
        super().setUp()
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Recepción'))
        self.client.force_login(self.tutor)

    def test_no_muestra_citas_pasadas_ni_cerradas(self):
        self.crear_cita(dias=-2)
        self.crear_cita(dias=3, estado='cancelada')
        self.crear_cita(dias=3)
        self.assertEqual(self.client.get('/riesgo/').context['resumen']['total'], 1)

    def test_el_filtro_por_nivel_no_cambia_el_resumen(self):
        self.crear_cita(dias=3)
        self.crear_cita(dias=4)
        respuesta = self.client.get('/riesgo/?nivel=alto')
        resumen = respuesta.context['resumen']
        self.assertEqual(resumen['total'], 2)
        self.assertEqual(respuesta.context['pagina'].paginator.count, resumen['alto'])

    def test_exporta_las_horas_en_riesgo(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.crear_cita(dias=3)
        respuesta = self.client.get('/riesgo/exportar/')
        hoja = load_workbook(BytesIO(respuesta.content)).active
        self.assertEqual(hoja['E2'].value, 'Bruno')


class ContactoSegunRiesgoTest(BaseCitas):
    """El mensaje de WhatsApp, el guion de llamada y el sobrecupo cambian según el riesgo."""

    def stats(self, **cambios):
        base = {'tutor': {}, 'dia': {}, 'tipo': {}, 'global': .2}
        base.update(cambios)
        return base

    def test_el_mensaje_cambia_segun_el_nivel(self):
        from .contacto import mensaje_segun_riesgo
        cita = self.crear_cita()
        alto = mensaje_segun_riesgo(cita, 'alto')
        medio = mensaje_segun_riesgo(cita, 'medio')
        bajo = mensaje_segun_riesgo(cita, 'bajo')
        self.assertIn('respondes SÍ', alto)
        self.assertIn('cambiarla', medio)
        self.assertNotIn('SÍ', bajo)
        self.assertTrue(all(m.startswith('Hola Ana') for m in (alto, medio, bajo)))

    def test_el_enlace_no_se_arma_sin_un_telefono_valido(self):
        from .contacto import enlace_whatsapp
        self.assertIsNone(enlace_whatsapp('', 'Hola'))
        self.assertIsNone(enlace_whatsapp('1234', 'Hola'))
        self.assertTrue(enlace_whatsapp('9 1234 5678', 'Hola').startswith('https://wa.me/56912345678'))

    def test_el_guion_menciona_las_faltas_del_tutor(self):
        from .contacto import guion_llamada
        cita = self.crear_cita()
        guion = guion_llamada(cita, 'alto', self.stats(tutor={self.cliente.pk: [3, 5]}))
        self.assertTrue(any('Ha faltado 3 de 5' in g for g in guion))
        self.assertTrue(any('sobrecupo' in g for g in guion))

    def test_el_guion_avisa_la_primera_visita_y_la_reserva_lejana(self):
        from .contacto import guion_llamada
        cita = self.crear_cita(dias=20, anticipacion=20)
        guion = guion_llamada(cita, 'medio', self.stats())
        self.assertTrue(any('primera visita' in g for g in guion))
        self.assertTrue(any('hace 20 días' in g for g in guion))

    def test_sin_factores_basta_un_recordatorio(self):
        from .contacto import guion_llamada
        cita = self.crear_cita(dias=1, anticipacion=1)
        guion = guion_llamada(cita, 'bajo', self.stats(tutor={self.cliente.pk: [0, 6]}))
        self.assertEqual(guion, ['Basta con un recordatorio breve.'])

    def test_el_panel_sugiere_sobrecupo_solo_para_riesgo_alto(self):
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Recepción'))
        self.client.force_login(self.tutor)
        self.crear_cita(dias=1)
        contexto = self.client.get('/').context
        altos = [c for c in contexto['confirmar'] if c['nivel'] == 'alto']
        self.assertEqual(len(contexto['sobrecupo']), len(altos))
        self.assertTrue(all(c['guion'] for c in contexto['confirmar']))


class ExportacionSeguraTest(BaseCitas):
    """La exportación no ejecuta fórmulas y no la usa quien solo ve citas."""

    def test_un_texto_con_igual_queda_como_texto(self):
        from io import BytesIO
        from django.contrib.auth.models import Group
        from openpyxl import load_workbook
        self.tutor.groups.add(Group.objects.get(name='Administración'))
        self.client.force_login(self.tutor)
        cita = self.crear_cita(estado='cancelada')
        Cita.objects.filter(pk=cita.pk).update(motivo_cancelacion='=1+1')
        hoja = load_workbook(BytesIO(self.client.get('/citas/exportar/').content)).active
        self.assertEqual(hoja['L2'].data_type, 's')
        self.assertEqual(hoja['L2'].value, '=1+1')

    def test_veterinarios_no_exportan(self):
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Veterinarios'))
        self.client.force_login(self.tutor)
        self.assertEqual(self.client.get('/citas/exportar/').status_code, 403)


class MantencionTest(TestCase):
    """CRUD de configuración y edición de citas, clientes y mascotas."""

    def setUp(self):
        from django.contrib.auth.models import Group

        from usuarios.models import Perfil
        self.admin = User.objects.create_user('admin1')
        self.admin.groups.add(Group.objects.get(name='Administración'))
        self.client.force_login(self.admin)
        self.vet = User.objects.create_user('vet1', first_name='Paz')
        Perfil.objects.update_or_create(usuario=self.vet, defaults={'rol': 'veterinario'})

    def test_crear_editar_y_eliminar_tipo_de_atencion(self):
        self.client.post('/citas/tipos/nuevo/', {'nombre': 'Ecografía', 'duracion_minutos': 40,
                                                  'precio': 35000, 'activo': 'on'})
        tipo = TipoAtencion.objects.get(nombre='Ecografía')
        self.client.post(f'/citas/tipos/{tipo.pk}/', {'nombre': 'Ecografía abdominal', 'duracion_minutos': 45,
                                                       'precio': 38000, 'activo': 'on'})
        tipo.refresh_from_db()
        self.assertEqual((tipo.nombre, tipo.precio), ('Ecografía abdominal', 38000))
        self.client.post(f'/citas/tipos/{tipo.pk}/eliminar/')
        self.assertFalse(TipoAtencion.objects.filter(pk=tipo.pk).exists())

    def test_no_elimina_un_tipo_con_citas(self):
        from mascotas.models import Cliente, Especie, Mascota
        tipo = TipoAtencion.objects.create(nombre='Control', duracion_minutos=30, precio=20000)
        cliente = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        mascota = Mascota.objects.create(nombre='Bruno', especie=Especie.objects.create(nombre='Perro'), sexo='macho')
        Cita.objects.create(veterinario=self.vet, mascota=mascota, cliente=cliente, tipo_atencion=tipo,
                            fecha=timezone.localdate(), hora='10:00')
        self.client.post(f'/citas/tipos/{tipo.pk}/eliminar/')
        self.assertTrue(TipoAtencion.objects.filter(pk=tipo.pk).exists())

    def test_los_horarios_no_se_pueden_cruzar(self):
        from .models import HorarioAtencion
        datos = {'veterinario': self.vet.pk, 'dia_semana': 0, 'hora_inicio': '09:00', 'hora_fin': '13:00', 'activo': 'on'}
        self.client.post('/citas/horarios/nuevo/', datos)
        respuesta = self.client.post('/citas/horarios/nuevo/', dict(datos, hora_inicio='12:00', hora_fin='14:00'))
        self.assertEqual(HorarioAtencion.objects.count(), 1)
        self.assertContains(respuesta, 'Se cruza con otro tramo')

    def test_la_hora_de_termino_debe_ser_posterior(self):
        respuesta = self.client.post('/citas/horarios/nuevo/', {'veterinario': self.vet.pk, 'dia_semana': 1,
                                                                  'hora_inicio': '15:00', 'hora_fin': '10:00', 'activo': 'on'})
        self.assertContains(respuesta, 'posterior a la de inicio')

    def test_editar_y_desactivar_cliente(self):
        from mascotas.models import Cliente
        cliente = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        self.client.post(f'/clientes/{cliente.pk}/editar/', {'rut': '15.234.871-1', 'nombres': 'Ana María',
                                                             'apellidos': 'Soto', 'telefono': '912345678'})
        cliente.refresh_from_db()
        self.assertEqual(cliente.nombres, 'Ana María')
        self.client.post(f'/clientes/{cliente.pk}/estado/')
        cliente.refresh_from_db()
        self.assertFalse(cliente.activo)

    def test_editar_la_cita_pendiente(self):
        from mascotas.models import Cliente, Especie, Mascota
        control = TipoAtencion.objects.create(nombre='Control', duracion_minutos=30, precio=20000)
        vacuna = TipoAtencion.objects.create(nombre='Vacunación', duracion_minutos=20, precio=15000)
        cliente = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        mascota = Mascota.objects.create(nombre='Bruno', especie=Especie.objects.create(nombre='Perro'), sexo='macho')
        Tenencia.objects.create(cliente=cliente, mascota=mascota, es_principal=True)
        cita = Cita.objects.create(veterinario=self.vet, mascota=mascota, cliente=cliente, tipo_atencion=control,
                                   fecha=timezone.localdate(), hora='10:00')
        self.client.post(f'/citas/{cita.pk}/editar/', {'tipo_atencion': vacuna.pk, 'notas': 'Refuerzo'})
        cita.refresh_from_db()
        self.assertEqual((cita.tipo_atencion, cita.notas), (vacuna, 'Refuerzo'))

    def test_no_se_elimina_un_grupo_base(self):
        from django.contrib.auth.models import Group
        grupo = Group.objects.get(name='Recepción')
        self.client.post(f'/administracion/grupos/{grupo.pk}/eliminar/')
        self.assertTrue(Group.objects.filter(pk=grupo.pk).exists())

    def test_las_pantallas_de_configuracion_abren(self):
        for url in ('/citas/tipos/', '/citas/tipos/nuevo/', '/citas/horarios/', '/citas/horarios/nuevo/',
                    '/mascotas/tipos/', '/mascotas/tipos/nuevo/', '/mascotas/razas/', '/mascotas/razas/nueva/'):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_recepcion_no_entra_a_configuracion(self):
        from django.contrib.auth.models import Group
        recep = User.objects.create_user('recep9')
        recep.groups.add(Group.objects.get(name='Recepción'))
        self.client.force_login(recep)
        self.assertEqual(self.client.get('/citas/tipos/nuevo/').status_code, 403)
        self.assertEqual(self.client.get('/citas/horarios/nuevo/').status_code, 403)

    def test_eliminar_un_tramo_de_horario(self):
        from .models import HorarioAtencion
        tramo = HorarioAtencion.objects.create(veterinario=self.vet, dia_semana=2, hora_inicio='09:00', hora_fin='12:00')
        self.client.post(f'/citas/horarios/{tramo.pk}/eliminar/')
        self.assertFalse(HorarioAtencion.objects.filter(pk=tramo.pk).exists())

    def test_editar_y_desactivar_mascota(self):
        from mascotas.models import Especie, Mascota
        canino = Especie.objects.get(nombre='Canino')
        bruno = Mascota.objects.create(nombre='Bruno', especie=canino, sexo='macho')
        self.client.post(f'/mascotas/{bruno.pk}/editar/', {'especie': canino.pk, 'nombre': 'Bruno II', 'sexo': 'macho',
                                                            'tiene_microchip': 'False', 'esterilizada': 'unknown'})
        bruno.refresh_from_db()
        self.assertEqual(bruno.nombre, 'Bruno II')
        self.client.post(f'/mascotas/{bruno.pk}/estado/')
        bruno.refresh_from_db()
        self.assertFalse(bruno.activo)
        self.client.post(f'/mascotas/{bruno.pk}/estado/')
        bruno.refresh_from_db()
        self.assertTrue(bruno.activo)

    def test_eliminar_tipo_de_mascota_y_raza_solo_si_no_se_usan(self):
        from mascotas.models import Especie, Mascota, Raza
        huron = Especie.objects.create(nombre='Hurón doméstico')
        raza = Raza.objects.create(especie=huron, nombre='Albino')
        Mascota.objects.create(nombre='Pelusa', especie=huron, raza=raza, sexo='hembra')
        self.client.post(f'/mascotas/razas/{raza.pk}/eliminar/')
        self.client.post(f'/mascotas/tipos/{huron.pk}/eliminar/')
        self.assertTrue(Raza.objects.filter(pk=raza.pk).exists())      # en uso: no se borra
        self.assertTrue(Especie.objects.filter(pk=huron.pk).exists())
        libre = Raza.objects.create(especie=huron, nombre='Canela')
        self.client.post(f'/mascotas/razas/{libre.pk}/eliminar/')
        self.assertFalse(Raza.objects.filter(pk=libre.pk).exists())
        vacio = Especie.objects.create(nombre='Pez dorado')
        self.client.post(f'/mascotas/tipos/{vacio.pk}/eliminar/')
        self.assertFalse(Especie.objects.filter(pk=vacio.pk).exists())


class AtencionClinicaTest(TestCase):
    """PV-FIC-01: el veterinario registra la atención al cerrar la cita."""

    def setUp(self):
        from django.contrib.auth.models import Group
        from usuarios.models import Perfil

        self.vet = User.objects.create_user('vet', first_name='Paz', last_name='Herrera')
        Perfil.objects.create(usuario=self.vet, rol='veterinario')
        self.vet.groups.add(Group.objects.get(name='Veterinarios'))
        canino = Especie.objects.create(nombre='Perro')
        self.ana = Cliente.objects.create(rut='15234871-1', nombres='Ana', apellidos='Soto')
        self.bruno = Mascota.objects.create(nombre='Bruno', especie=canino, sexo='macho')
        Tenencia.objects.create(cliente=self.ana, mascota=self.bruno, es_principal=True)
        self.control = TipoAtencion.objects.create(nombre='Control', duracion_minutos=30, precio=20000)
        self.client.force_login(self.vet)

    def cita(self, dias=0, estado='pendiente', hora='10:00'):
        return Cita.objects.create(
            veterinario=self.vet, mascota=self.bruno, cliente=self.ana, tipo_atencion=self.control,
            fecha=timezone.localdate() + timedelta(days=dias), hora=hora, estado=estado, notas='Vómitos')

    def datos(self, **extra):
        return {'motivo': 'Vómitos', 'diagnostico': 'Gastritis leve', 'tratamiento': 'Dieta blanda 3 días',
                'peso_kg': '12.40', **extra}

    def test_el_formulario_parte_con_el_motivo_de_la_cita(self):
        c = self.cita()
        respuesta = self.client.get(f'/citas/{c.pk}/atender/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['form'].initial['motivo'], 'Vómitos')

    def test_registrar_cierra_la_cita_y_guarda_la_atencion(self):
        from .models import Atencion
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/atender/', self.datos())
        c.refresh_from_db()
        self.assertEqual(c.estado, 'completada')
        atencion = Atencion.objects.get(cita=c)
        self.assertEqual(atencion.diagnostico, 'Gastritis leve')
        self.assertEqual(atencion.registrada_por, self.vet)

    def test_el_peso_se_actualiza_en_la_ficha_de_la_mascota(self):
        from decimal import Decimal
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/atender/', self.datos(peso_kg='13.05'))
        self.bruno.refresh_from_db()
        self.assertEqual(self.bruno.peso_kg, Decimal('13.05'))

    def test_sin_diagnostico_no_se_cierra(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/atender/', self.datos(diagnostico=''))
        c.refresh_from_db()
        self.assertEqual(c.estado, 'pendiente')

    def test_peso_imposible(self):
        c = self.cita()
        respuesta = self.client.post(f'/citas/{c.pk}/atender/', self.datos(peso_kg='-3'))
        self.assertIn('peso_kg', respuesta.context['form'].errors)

    def test_no_se_atiende_una_cita_futura(self):
        c = self.cita(dias=3)
        self.client.post(f'/citas/{c.pk}/atender/', self.datos())
        c.refresh_from_db()
        self.assertEqual(c.estado, 'pendiente')

    def test_una_cita_ya_cerrada_no_se_vuelve_a_atender(self):
        from .models import Atencion
        c = self.cita(estado='no_asistio')
        self.client.post(f'/citas/{c.pk}/atender/', self.datos())
        self.assertFalse(Atencion.objects.exists())

    def test_vuelve_solo_a_direcciones_del_sistema(self):
        c = self.cita()
        respuesta = self.client.post(f'/citas/{c.pk}/atender/', self.datos(next='https://otro-sitio.cl/'))
        self.assertTrue(respuesta['Location'].startswith('/citas/calendario/'))


class ConfirmacionTest(CalendarioTest):
    """PV-CIT-05: estado de confirmación con fecha y usuario."""

    def test_registrar_guarda_resultado_fecha_y_usuario(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'confirmada'})
        c.refresh_from_db()
        self.assertEqual(c.confirmacion, 'confirmada')
        self.assertEqual(c.confirmacion_por, self.recepcion)
        self.assertIsNotNone(c.confirmacion_fecha)

    def test_un_resultado_inventado_no_se_guarda(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'quizas'})
        c.refresh_from_db()
        self.assertEqual(c.confirmacion, '')

    def test_pidio_cambio_lleva_a_reagendar(self):
        c = self.cita()
        respuesta = self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'pidio_cambio'})
        self.assertEqual(respuesta['Location'], f'/citas/{c.pk}/reagendar/')

    def test_no_se_confirma_una_cita_cerrada(self):
        c = self.cita(estado='cancelada')
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'confirmada'})
        c.refresh_from_db()
        self.assertEqual(c.confirmacion, '')

    def test_se_ve_en_el_calendario_y_en_el_detalle(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'confirmada'})
        evento = self.eventos()[0]
        self.assertIn('conf-confirmada', evento['classNames'])
        self.assertEqual(evento['extendedProps']['confirmacion'], 'Confirmó')
        self.assertTrue(evento['title'].startswith('✓'))
        self.assertContains(self.client.get(f'/citas/{c.pk}/detalle/'), 'Confirmó')

    def test_se_ve_en_horas_en_riesgo(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'no_contesta'})
        self.assertContains(self.client.get('/riesgo/'), 'No contesta')

    def test_al_reagendar_hay_que_volver_a_confirmar(self):
        c = self.cita()
        self.client.post(f'/citas/{c.pk}/confirmacion/', {'resultado': 'confirmada'})
        self.client.post(f'/citas/{c.pk}/reagendar/', {
            'veterinario': self.vet.pk, 'fecha': self.manana.isoformat(), 'hora': '11:00'})
        c.refresh_from_db()
        self.assertEqual(c.hora.strftime('%H:%M'), '11:00')
        self.assertEqual(c.confirmacion, '')
        self.assertIsNone(c.confirmacion_por)


class ListaLlamadasTest(BaseCitas):
    """PV-REC-02: lista de llamadas del día siguiente ordenada por riesgo."""

    def setUp(self):
        super().setUp()
        from django.contrib.auth.models import Group
        self.tutor.groups.add(Group.objects.get(name='Recepción'))
        self.client.force_login(self.tutor)
        # Otro tutor con historial de faltas: sus citas salen con más riesgo.
        self.faltador = Cliente.objects.create(rut='16890342-1', nombres='Matías', apellidos='Silva',
                                               telefono='+56 9 8765 4321')
        self.luna = Mascota.objects.create(nombre='Luna', especie=self.especie, sexo='hembra')
        Tenencia.objects.create(cliente=self.faltador, mascota=self.luna, es_principal=True)
        for dias in (-30, -20, -10):
            self.crear_cita(dias=dias, mascota=self.luna, estado='no_asistio')
        self.crear_cita(dias=-15, estado='completada')

    def test_ordena_las_de_manana_de_mayor_a_menor_riesgo(self):
        tranquila = self.crear_cita(dias=1)
        riesgosa = self.crear_cita(dias=1, mascota=self.luna)
        self.crear_cita(dias=2)                      # pasado mañana no aparece
        filas = self.client.get('/citas/llamadas/').context['filas']
        self.assertEqual([f['cita'] for f in filas], [riesgosa, tranquila])
        self.assertGreater(filas[0]['puntaje'], filas[1]['puntaje'])

    def test_trae_el_whatsapp_con_el_mensaje(self):
        self.crear_cita(dias=1, mascota=self.luna)
        fila = self.client.get('/citas/llamadas/').context['filas'][0]
        self.assertTrue(fila['whatsapp'].startswith('https://wa.me/56987654321'))

    def test_registrar_el_resultado_vuelve_a_la_lista(self):
        cita = self.crear_cita(dias=1)
        respuesta = self.client.post(f'/citas/{cita.pk}/confirmacion/',
                                     {'resultado': 'no_contesta', 'next': '/citas/llamadas/'})
        self.assertEqual(respuesta['Location'], '/citas/llamadas/')
        avance = self.client.get('/citas/llamadas/').context['avance']
        self.assertEqual((avance['total'], avance['no_contesta'], avance['contactadas']), (1, 1, 0))

    def test_solo_por_llamar_esconde_las_ya_confirmadas(self):
        confirmada = self.crear_cita(dias=1)
        sin_respuesta = self.crear_cita(dias=1, mascota=self.luna)
        self.client.post(f'/citas/{confirmada.pk}/confirmacion/', {'resultado': 'confirmada'})
        self.client.post(f'/citas/{sin_respuesta.pk}/confirmacion/', {'resultado': 'no_contesta'})
        filas = self.client.get('/citas/llamadas/', {'filtro': 'pendientes'}).context['filas']
        self.assertEqual([f['cita'] for f in filas], [sin_respuesta])

    def test_otra_fecha_y_fechas_pasadas(self):
        cita = self.crear_cita(dias=2)
        contexto = self.client.get('/citas/llamadas/', {'fecha': cita.fecha.isoformat()}).context
        self.assertEqual(len(contexto['filas']), 1)
        ayer = (timezone.localdate() - timedelta(days=1)).isoformat()
        self.assertEqual(self.client.get('/citas/llamadas/', {'fecha': ayer}).context['fecha'],
                         timezone.localdate())

    def test_sin_permiso_de_riesgo_no_entra(self):
        from django.contrib.auth.models import Group
        self.tutor.groups.set([Group.objects.get(name='Veterinarios')])
        self.assertEqual(self.client.get('/citas/llamadas/').status_code, 403)


class EtapasDeAtencionTest(PanelDelDiaTest):
    """PV-REC-05: etapas "En atención" y "Lista para salir" en el flujo de pacientes."""

    def test_pasar_a_atencion(self):
        c = self.crear_cita(dias=0)
        self.client.post(f'/citas/{c.pk}/llego/')
        self.client.post(f'/citas/{c.pk}/pasar/')
        c.refresh_from_db()
        self.assertIsNotNone(c.inicio_atencion)
        contexto = self.client.get('/').context
        self.assertEqual([t['cita'] for t in contexto['en_atencion']], [c])
        self.assertEqual(contexto['en_sala'], [])

    def test_pasar_sin_marcar_llegada_registra_la_llegada(self):
        c = self.crear_cita(dias=0)
        self.client.post(f'/citas/{c.pk}/pasar/')
        c.refresh_from_db()
        self.assertIsNotNone(c.hora_llegada)

    def test_solo_las_citas_de_hoy_pasan_a_atencion(self):
        c = self.crear_cita(dias=1)
        self.client.post(f'/citas/{c.pk}/pasar/')
        c.refresh_from_db()
        self.assertIsNone(c.inicio_atencion)

    def test_atendida_queda_lista_para_salir_hasta_que_sale(self):
        c = self.crear_cita(dias=0, estado='completada')
        contexto = self.client.get('/').context
        self.assertEqual([t['cita'] for t in contexto['listas']], [c])
        self.client.post(f'/citas/{c.pk}/salio/')
        c.refresh_from_db()
        self.assertIsNotNone(c.hora_salida)
        contexto = self.client.get('/').context
        self.assertEqual(contexto['listas'], [])
        self.assertEqual(len(contexto['atendidas']), 1)     # sigue contando como atendida

    def test_salio_solo_para_citas_atendidas(self):
        c = self.crear_cita(dias=0)
        self.client.post(f'/citas/{c.pk}/salio/')
        c.refresh_from_db()
        self.assertIsNone(c.hora_salida)
