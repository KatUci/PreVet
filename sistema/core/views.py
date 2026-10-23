from django.contrib.auth.decorators import login_required, permission_required
from django.db.models import Q, Sum
from django.shortcuts import render
from django.utils import timezone

from citas.models import Cita


@login_required
def inicio(request):
    """Tablero con los indicadores de inasistencia."""

    # Es la página a la que llega todo el mundo al iniciar sesión. Quien
    # todavía no tiene permisos ve un aviso en vez de un error.
    # Los clientes con cuenta tienen su propio portal.
    if getattr(request.user, 'cliente', None) is not None and not request.user.has_perm('citas.view_cita'):
        from django.shortcuts import redirect
        return redirect('portal')

    if not request.user.has_perm('citas.view_cita'):
        return render(request, 'core/sin_permisos.html', {'seccion': 'inicio'})

    # Recepción y veterinarios trabajan con el día a día: su inicio es el
    # panel "Hoy en la clínica". Los indicadores son para administración.
    if not request.user.has_perm('citas.ver_indicadores'):
        from citas.panel import panel_dia
        return panel_dia(request)

    hoy = timezone.localdate()

    # Solo las citas que ya ocurrieron: de las futuras todavía
    # no se sabe si van a concretarse.
    ocurridas = Cita.objects.filter(fecha__lt=hoy).exclude(estado='pendiente')

    total = ocurridas.count()
    faltas = ocurridas.filter(estado='no_asistio')
    n_faltas = faltas.count()

    tasa = (n_faltas / total * 100) if total else 0
    perdido = faltas.aggregate(t=Sum('tipo_atencion__precio'))['t'] or 0

    # Inasistencias por día de la semana, para ver si hay patrón.
    nombres = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
    por_dia = []
    for numero, nombre in enumerate(nombres):
        del_dia = [c for c in ocurridas if c.fecha.weekday() == numero]
        n = len(del_dia)
        f = len([c for c in del_dia if c.estado == 'no_asistio'])
        if n:
            por_dia.append({
                'dia': nombre,
                'total': n,
                'faltas': f,
                'tasa': round(f / n * 100),
            })

    tope = max([d['tasa'] for d in por_dia] + [1])
    for d in por_dia:
        d['alto'] = round(d['tasa'] / tope * 100)

    # ---------- Últimos 30 días contra los 30 anteriores ----------
    import datetime

    def periodo(desde, hasta):
        citas = [c for c in ocurridas if desde <= c.fecha <= hasta]
        f = [c for c in citas if c.estado == 'no_asistio']
        monto = sum(c.tipo_atencion.precio for c in f)
        return {'total': len(citas), 'faltas': len(f),
                'tasa': round(len(f) / len(citas) * 100, 1) if citas else 0, 'perdido': monto}

    ocurridas = list(ocurridas.select_related('tipo_atencion'))
    actual = periodo(hoy - datetime.timedelta(days=30), hoy - datetime.timedelta(days=1))
    anterior = periodo(hoy - datetime.timedelta(days=60), hoy - datetime.timedelta(days=31))
    variacion = {}
    for k in ('tasa', 'faltas', 'total'):
        diferencia = round(actual[k] - anterior[k], 1)
        variacion[k] = {'valor': abs(diferencia),
                        'signo': 'sube' if diferencia > 0 else 'baja' if diferencia < 0 else 'igual'}
    actual['perdido_txt'] = f"{actual['perdido']:,}".replace(',', '.')

    # ---------- Evolución mensual (últimos 6 meses con datos) ----------
    meses = {}
    for c in ocurridas:
        meses.setdefault((c.fecha.year, c.fecha.month), []).append(c)
    nombres_mes = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
    mensual = []
    for (anio, mes) in sorted(meses)[-6:]:
        grupo = meses[(anio, mes)]
        f = sum(1 for c in grupo if c.estado == 'no_asistio')
        mensual.append({'mes': f'{nombres_mes[mes - 1]} {str(anio)[2:]}', 'total': len(grupo), 'faltas': f,
                        'tasa': round(f / len(grupo) * 100, 1)})
    tope_mes = max([m['tasa'] for m in mensual] + [1])
    for m in mensual:
        m['alto'] = round(m['tasa'] / tope_mes * 100)

    # ---------- Hoy y las citas de mayor riesgo de la semana ----------
    de_hoy = Cita.objects.filter(fecha=hoy).exclude(estado='cancelada')
    hoy_resumen = {
        'total': de_hoy.count(),
        'por_llegar': de_hoy.filter(estado='pendiente', hora_llegada__isnull=True).count(),
        'en_sala': de_hoy.filter(estado='pendiente', hora_llegada__isnull=False,
                                 inicio_atencion__isnull=True).count(),
        'en_atencion': de_hoy.filter(estado='pendiente', inicio_atencion__isnull=False).count(),
        'atendidas': de_hoy.filter(estado='completada').count(),
    }
    from citas.riesgo import estadisticas, puntuar
    semana = (Cita.objects.filter(fecha__gte=hoy, fecha__lte=hoy + datetime.timedelta(days=6), estado='pendiente')
              .select_related('mascota', 'cliente', 'veterinario', 'tipo_atencion'))
    stats = estadisticas()
    en_riesgo = []
    for cita in semana:
        puntaje, nivel, _, _ = puntuar(cita, stats)
        if nivel == 'alto':
            en_riesgo.append({'cita': cita, 'puntaje': puntaje})
    en_riesgo.sort(key=lambda e: -e['puntaje'])

    contexto = {
        'seccion': 'inicio',
        'actual': actual,
        'variacion': variacion,
        'mensual': mensual,
        'hoy_resumen': hoy_resumen,
        'en_riesgo': en_riesgo[:5],
        'en_riesgo_total': len(en_riesgo),
        'total_ocurridas': total,
        'total_faltas': n_faltas,
        'tasa': round(tasa, 1),
        'perdido': perdido,
        # Con punto de miles, como se escriben los pesos en Chile.
        'perdido_txt': f"{perdido:,}".replace(',', '.'),
        'por_dia': por_dia,
        'proximas': Cita.objects.filter(fecha__gte=hoy, estado='pendiente')
                                .select_related('mascota', 'veterinario', 'tipo_atencion')
                                .order_by('fecha', 'hora')[:8],
    }
    return render(request, 'core/inicio.html', contexto)

NIVELES = (('alto', 'Alto'), ('medio', 'Medio'), ('bajo', 'Bajo'))


def _evaluar_riesgo(g):
    """Citas pendientes desde hoy con su riesgo, según los filtros de la dirección.

    Devuelve la lista evaluada (ya filtrada por nivel y ordenada de mayor a
    menor riesgo), los filtros aplicados y el resumen por nivel antes de
    filtrar por nivel, para que las tarjetas muestren el panorama completo.
    """
    from citas.calendario import _enlace_whatsapp
    from citas.contacto import guion_llamada
    from citas.riesgo import estadisticas, puntuar
    from citas.views import _filtrar_citas

    hoy = timezone.localdate()
    desde = g.get('desde', '')
    if not desde or desde < hoy.isoformat():
        desde = hoy.isoformat()
    citas, filtros = _filtrar_citas({
        'q': g.get('q', ''), 'estado': 'pendiente', 'veterinario': g.get('veterinario', ''),
        'desde': desde, 'hasta': g.get('hasta', ''), 'orden': 'antiguas',
    })
    filtros['nivel'] = g.get('nivel', '') if g.get('nivel', '') in dict(NIVELES) else ''

    stats = estadisticas()
    evaluadas = []
    for cita in citas:
        puntaje, nivel, accion, detalle = puntuar(cita, stats)
        evaluadas.append({
            'cita': cita,
            'puntaje': puntaje,
            'nivel': nivel,
            'accion': accion,
            'detalle': detalle,
            'whatsapp': _enlace_whatsapp(cita, nivel),
            'guion': guion_llamada(cita, nivel, stats),
        })

    resumen = {n: sum(1 for e in evaluadas if e['nivel'] == n) for n, _ in NIVELES}
    resumen['total'] = len(evaluadas)
    resumen['monto'] = sum(e['cita'].tipo_atencion.precio for e in evaluadas if e['nivel'] != 'bajo')
    resumen['monto_txt'] = f"{resumen['monto']:,}".replace(',', '.')
    for n, _ in NIVELES:
        resumen[f'pct_{n}'] = round(resumen[n] / resumen['total'] * 100) if resumen['total'] else 0

    if filtros['nivel']:
        evaluadas = [e for e in evaluadas if e['nivel'] == filtros['nivel']]
    evaluadas.sort(key=lambda e: (-e['puntaje'], e['cita'].fecha, e['cita'].hora))
    return evaluadas, filtros, resumen


@login_required
@permission_required('citas.ver_riesgo', raise_exception=True)
def riesgo(request):
    """Citas pendientes ordenadas por su riesgo de inasistencia."""

    import datetime

    from django.contrib.auth.models import User
    from django.core.paginator import Paginator

    from citas import riesgo as modelo

    evaluadas, filtros, resumen = _evaluar_riesgo(request.GET)
    pagina = Paginator(evaluadas, 15).get_page(request.GET.get('pagina'))

    parametros = request.GET.copy()
    parametros.pop('pagina', None)
    sin_nivel = parametros.copy()
    sin_nivel.pop('nivel', None)

    hoy = timezone.localdate()
    manana = hoy + datetime.timedelta(days=1)
    return render(request, 'core/riesgo.html', {
        'seccion': 'riesgo',
        'pagina': pagina,
        'rango_paginas': pagina.paginator.get_elided_page_range(pagina.number, on_each_side=2, on_ends=1),
        'filtros': filtros,
        'resumen': resumen,
        'niveles': NIVELES,
        'qs': parametros.urlencode(),
        'qs_sin_nivel': sin_nivel.urlencode(),
        'veterinarios': User.objects.filter(perfil__rol='veterinario').order_by('first_name'),
        'atajos': [('Hoy', hoy, hoy), ('Mañana', manana, manana),
                   ('Próximos 7 días', hoy, hoy + datetime.timedelta(days=6))],
        'hay_filtros': any(request.GET.get(k) for k in ('q', 'nivel', 'veterinario', 'desde', 'hasta')),
        'pesos': [('Historial de faltas del tutor', modelo.PESO_TUTOR),
                  ('Días de anticipación de la reserva', modelo.PESO_ANTICIPACION),
                  ('Día de la semana', modelo.PESO_DIA),
                  ('Tipo de atención', modelo.PESO_TIPO)],
        'umbral_alto': modelo.UMBRAL_ALTO,
        'umbral_medio': modelo.UMBRAL_MEDIO,
        'acciones': modelo.ACCIONES,
    })


@login_required
@permission_required('citas.ver_riesgo', raise_exception=True)
def riesgo_exportar(request):
    """Descarga en Excel las horas en riesgo, con los mismos filtros de la pantalla."""

    from django.http import HttpResponse
    from openpyxl import Workbook

    from citas.views import texto_seguro
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    evaluadas, _, _ = _evaluar_riesgo(request.GET)

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Horas en riesgo'
    columnas = [('Puntaje', 9), ('Nivel', 8), ('Fecha', 12), ('Hora', 8), ('Mascota', 16), ('Cliente', 26),
                ('Teléfono', 15), ('Profesional', 24), ('Atención', 22), ('Acción sugerida', 46),
                ('Factores', 70)]
    hoja.append([c for c, _ in columnas])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='2B2160')
    for i, (_, ancho) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for e in evaluadas:
        c = e['cita']
        hoja.append([e['puntaje'], e['nivel'].capitalize(), c.fecha, c.hora.strftime('%H:%M'), c.mascota.nombre,
                     c.cliente.nombre_completo, c.cliente.telefono, c.veterinario.get_full_name(),
                     c.tipo_atencion.nombre, e['accion'], ' · '.join(e['detalle'])])
    for fila in hoja.iter_rows(min_row=2):
        fila[2].number_format = 'DD-MM-YYYY'
    texto_seguro(hoja)
    hoja.freeze_panes = 'A2'
    hoja.auto_filter.ref = hoja.dimensions

    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="horas_en_riesgo_{timezone.localdate():%Y-%m-%d}.xlsx"'
    libro.save(respuesta)
    return respuesta


@login_required
@permission_required('citas.ver_indicadores', raise_exception=True)
def indicadores(request):
    """Análisis de las inasistencias, factor por factor del modelo de riesgo."""

    hoy = timezone.localdate()

    # Se traen todas de una vez y el resto del cálculo se hace en Python.
    # Son poco más de 200 citas: cabe de sobra en memoria y evita decenas
    # de consultas separadas contra la base de datos.
    ocurridas = list(Cita.objects
                     .filter(fecha__lt=hoy)
                     .exclude(estado='pendiente')
                     .select_related('tipo_atencion', 'cliente'))

    def resumen(grupo):
        """Cuenta total, faltas, tasa y plata perdida de un grupo de citas."""
        faltas = [c for c in grupo if c.estado == 'no_asistio']
        perdido = sum(c.tipo_atencion.precio for c in faltas)
        return {
            'total': len(grupo),
            'faltas': len(faltas),
            'tasa': round(len(faltas) / len(grupo) * 100, 1) if grupo else 0,
            'perdido': perdido,
            # El monto se arma acá, con punto como separador de miles, que es
            # como se escriben los pesos en Chile. El filtro intcomma de Django
            # depende del idioma configurado y deja espacios en vez de puntos.
            'perdido_txt': f"{perdido:,}".replace(",", "."),
        }

    def escalar(filas):
        """Agrega el ancho de la barra, en porcentaje de la mayor."""
        tope = max([f['tasa'] for f in filas] + [1])
        for f in filas:
            f['ancho'] = round(f['tasa'] / tope * 100)
        return filas

    # ---------- Por tipo de atención (peso 20 en el modelo) ----------
    tipos = {}
    for c in ocurridas:
        tipos.setdefault(c.tipo_atencion.nombre, []).append(c)

    por_tipo = escalar(sorted(
        [dict(nombre=n, **resumen(g)) for n, g in tipos.items()],
        key=lambda f: -f['tasa']))

    # ---------- Por anticipación (peso 20 en el modelo) ----------
    TRAMOS = [(1, 3, '1 a 3 días'), (4, 7, '4 a 7 días'),
              (8, 14, '8 a 14 días'), (15, 99, '15 días o más')]

    por_anticipacion = escalar([
        dict(nombre=etiqueta,
             **resumen([c for c in ocurridas
                        if desde <= c.dias_anticipacion <= hasta]))
        for desde, hasta, etiqueta in TRAMOS])

    # ---------- Tutores reincidentes (peso 40, el más alto) ----------
    tutores = {}
    for c in ocurridas:
        tutores.setdefault(c.cliente, []).append(c)

    reincidentes = sorted(
        [dict(nombre=d.nombre_completo, **resumen(g))
         for d, g in tutores.items()],
        key=lambda f: (-f['faltas'], -f['tasa']))[:8]

    return render(request, 'core/indicadores.html', {
        'seccion': 'indicadores',
        'por_tipo': por_tipo,
        'por_anticipacion': por_anticipacion,
        'reincidentes': reincidentes,
        'general': resumen(ocurridas),
    })


@login_required
def buscar(request):
    """Buscador global de la barra superior: clientes, mascotas y citas."""

    from mascotas.directorio import filtrar_clientes
    from mascotas.models import Cliente, Mascota

    q = request.GET.get('q', '').strip()
    resultados = {'clientes': [], 'mascotas': [], 'citas': []}
    totales = {'clientes': 0, 'mascotas': 0, 'citas': 0}

    if len(q) >= 2:
        if request.user.has_perm('mascotas.view_cliente'):
            clientes = filtrar_clientes(Cliente.objects.filter(activo=True), q).order_by('apellidos', 'nombres')
            totales['clientes'] = clientes.count()
            resultados['clientes'] = clientes[:8]
        if request.user.has_perm('mascotas.view_mascota'):
            mascotas = Mascota.objects.select_related('especie', 'raza')
            for palabra in q.split():
                mascotas = mascotas.filter(Q(nombre__icontains=palabra) | Q(tutores__nombres__icontains=palabra)
                                           | Q(tutores__apellidos__icontains=palabra)
                                           | Q(tutores__rut__startswith=palabra.replace('.', '')))
            mascotas = mascotas.distinct().order_by('nombre')
            totales['mascotas'] = mascotas.count()
            resultados['mascotas'] = mascotas[:8]
        if request.user.has_perm('citas.view_cita'):
            from citas.views import _filtrar_citas
            citas, _ = _filtrar_citas({'q': q, 'orden': 'recientes'})
            totales['citas'] = citas.count()
            resultados['citas'] = citas[:8]

    return render(request, 'core/buscar.html', {
        'seccion': 'buscar',
        'q': q,
        'resultados': resultados,
        'totales': totales,
        'total': sum(totales.values()),
    })
